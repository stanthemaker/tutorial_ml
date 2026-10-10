"""Build the Week 5 slide deck: slides/week5_diffusion.pptx

    pip install python-pptx
    python slides/build_slides.py

Nothing on a slide is typed in by hand if it can be computed:

  * SHAPES and PARAMETER COUNTS come from the real UNet in part1_unet_diffusion.py
    (forward hooks on a real forward pass).
  * The NOISED DIGITS, the EMBEDDING heat map and the RECEPTIVE FIELDS are
    computed here from real MNIST digits and the real model code.
  * Part 2's SHAPES, PARAMETER COUNTS and GFLOPs come from the real DiT in
    part2_transformer_diffusion.py; its ATTENTION MAPS from its checkpoint.
  * EVIDENCE (loss curve, sample grids, timings) is read from results/, which
    part1_unet_diffusion.py, part2_transformer_diffusion.py and sample.py
    write (results/part1_unet/, results/part2_transformer/). If a file is missing, the slide shows the
    command that produces it instead of a made-up picture.

The slide kit (blank/header/bullets/...) is copied from week 4 so the decks
look the same.
"""

import copy
import json
import os
import sys
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import torch
import torch.nn as nn
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
WEEK5 = os.path.dirname(HERE)
if WEEK5 not in sys.path:                      # so this runs from any folder
    sys.path.insert(0, WEEK5)

from common import IMAGE_SIZE, NoiseSchedule, count_params, load_mnist, timestep_embedding
from part1_unet_diffusion import Block, UNet
from part2_transformer_diffusion import DiT

RESULTS = os.path.join(WEEK5, "results", "part1_unet")
CKPT = os.path.join(WEEK5, "checkpoints", "part1_unet.pt")
RESULTS2 = os.path.join(WEEK5, "results", "part2_transformer")
CKPT2 = os.path.join(WEEK5, "checkpoints", "part2_transformer.pt")
OUT = os.path.join(HERE, "week5_diffusion.pptx")

# ------------------------------------------------------------------ palette
# Identical to weeks 3 and 4: this is the next session of the same course, and
# a student should not have to re-learn what the colours mean.

INK = RGBColor(0x1C, 0x22, 0x2B)
MUTED = RGBColor(0x6B, 0x75, 0x82)
ACCENT = RGBColor(0xC2, 0x55, 0x1E)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
PAPER = RGBColor(0xFA, 0xF9, 0xF7)
RULE = RGBColor(0xD8, 0xD3, 0xCB)

C_WARM = "#c2551e"      # the new thing / the thing that wins
C_COOL = "#1f4e79"      # the baseline it is being compared against
C_GREY = "#6b7582"      # the oldest/weakest series, always also dashed
C_GRID = "#d8d3cb"
PAPER_HEX = "#faf9f7"

SANS = "Helvetica Neue"
MONO = "Menlo"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.85)


# ---------------------------------------------------------------- figures

def save(fig, path, pad=0.08):
    fig.savefig(path, dpi=200, bbox_inches="tight", pad_inches=pad,
                facecolor=PAPER_HEX)
    plt.close(fig)
    return path


def style(ax, xlabel="", ylabel=""):
    ax.set_facecolor(PAPER_HEX)
    ax.grid(True, color=C_GRID, lw=0.8, alpha=0.9)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(C_GRID)
    ax.set_xlabel(xlabel); ax.set_ylabel(ylabel)


# ------------------------------------------------------------------- maths
# python-pptx cannot write PowerPoint equations, so each formula is typeset
# with matplotlib mathtext and embedded as a transparent PNG at native size.

MATH_DPI = 300
_math_seq = [0]


def math_png(expr, path, fontsize=30, color="#1c222b"):
    fig = plt.figure(figsize=(0.1, 0.1))
    fig.text(0, 0, expr, fontsize=fontsize, color=color)
    fig.savefig(path, dpi=MATH_DPI, transparent=True, bbox_inches="tight",
                pad_inches=0.04)
    plt.close(fig)
    return path


def math_block(slide, expr, x, y, tmp, fontsize=30, color="#1c222b", center=False):
    _math_seq[0] += 1
    path = os.path.join(tmp, f"eq{_math_seq[0]}.png")
    math_png(expr, path, fontsize=fontsize, color=color)
    px_h, px_w = plt.imread(path).shape[:2]
    w = int(px_w / MATH_DPI * 914400)
    h = int(px_h / MATH_DPI * 914400)
    # A long equation rendered at its native size can be wider than the slide;
    # shrink it to fit rather than letting it hang over the edge.
    limit = int(W - 2 * M)
    if w > limit:
        h = int(h * limit / w)
        w = limit
    left = int((W - w) / 2) if center else x
    slide.shapes.add_picture(path, left, y, width=Emu(w), height=Emu(h))
    return Emu(w), Emu(h)


# ------------------------------------------------------------ slide kit


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    return s


def text(slide, x, y, w, h, runs_spec, align=PP_ALIGN.LEFT, spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (s, size, bold, color, font) in enumerate(runs_spec):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = s
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return box


def rule(slide, y, x=M, w=None):
    from pptx.enum.shapes import MSO_SHAPE
    w = w or (W - 2 * M)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(1.2))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def header(slide, kicker, title):
    if kicker:
        text(slide, M, Inches(0.5), W - 2 * M, Inches(0.35),
             [(kicker.upper(), 12, True, ACCENT, SANS)])
    text(slide, M, Inches(0.85), W - 2 * M, Inches(0.8),
         [(title, 29, True, BLUE, SANS)])
    rule(slide, Inches(1.62))


def bullets(slide, items, top=Inches(2.0), size=17, width=None, left=M):
    specs = []
    for it in items:
        kind = "body"
        if isinstance(it, tuple):
            it, kind = it
        if kind == "mono":
            specs.append((it, 15, False, INK, MONO))
        elif kind == "note":
            specs.append((it, 15, False, MUTED, SANS))
        elif kind == "lead":
            specs.append((it, size + 2, True, INK, SANS))
        else:
            specs.append((it, size, False, INK, SANS))
    return text(slide, left, top, width or (W - 2 * M), H - top - Inches(0.5),
                specs, spacing=1.15)


def picture(slide, path, top=Inches(1.95), height=Inches(4.3)):
    """Place a figure centred, scaled to `height` -- unless that would run off
    the sides, in which case width wins. A wide figure asked for by height is
    the easiest way to push a picture past the slide edge, and the overhang is
    invisible until someone opens the deck."""
    pic = slide.shapes.add_picture(path, Emu(0), top, height=height)
    limit = W - Inches(0.5)
    if pic.width > limit:
        pic.height = int(pic.height * limit / pic.width)
        pic.width = int(limit)
    pic.left = int((W - pic.width) / 2)
    return pic


def caption(slide, s, y=Inches(6.45)):
    text(slide, M, y, W - 2 * M, Inches(0.6), [(s, 14, False, MUTED, SANS)],
         align=PP_ALIGN.CENTER)


def content_slide(prs, kicker, title, items, top=Inches(2.0)):
    s = blank(prs)
    header(s, kicker, title)
    bullets(s, items, top=top)
    return s


def flow_slide(prs, kicker, title, blocks, tmp, top=Inches(1.95)):
    s = blank(prs)
    header(s, kicker, title)
    y = top
    for blk in blocks:
        kind = blk[0]
        if kind == "gap":
            y = Emu(int(y) + Inches(blk[1]))
            continue
        if kind == "eq":
            size = blk[2] if len(blk) > 2 else 28
            _, h = math_block(s, blk[1], M + Inches(0.4), y, tmp, fontsize=size)
            y = Emu(int(y) + int(h) + Inches(0.26))
            continue
        size, bold, color = {"lead": (18, True, INK), "text": (17, False, INK),
                             "note": (15, False, MUTED)}[kind]
        text(s, M, y, W - 2 * M, Inches(0.5), [(blk[1], size, bold, color, SANS)],
             spacing=1.15)
        chars_per_line = int((W - 2 * M) / Emu(1) / 914400 * 72 / (size * 0.50))
        lines = max(1, -(-len(blk[1]) // max(1, chars_per_line)))
        y = Emu(int(y) + Inches(lines * size * 1.15 / 72 + 0.16))
    return s


def code_slide(prs, kicker, title, source, lines, note=None, size=14.5):
    s = blank(prs)
    header(s, kicker, title)
    text(s, M, Inches(1.90), W - 2 * M, Inches(0.32),
         [(source, 13.5, True, MUTED, MONO)])
    specs = []
    for ln in lines:
        if isinstance(ln, tuple):
            specs.append(("  " + ln[0], size, True, ACCENT, MONO))
        else:
            specs.append(("  " + ln, size, False, INK, MONO))
    text(s, M, Inches(2.32), W - 2 * M, H - Inches(3.0), specs, spacing=1.08)
    if note:
        text(s, M, H - Inches(1.02), W - 2 * M, Inches(0.7),
             [(note, 14.5, False, MUTED, SANS)], spacing=1.15)
    return s


def table_slide(prs, kicker, title, headers, rows, note=None, widths=None,
                top=Inches(2.05)):
    s = blank(prs)
    header(s, kicker, title)
    ncol = len(headers)
    widths = widths or [1.0] * ncol
    scale = (W - 2 * M) / sum(widths)
    xs, x = [], M
    for w in widths:
        xs.append((x, int(w * scale)))
        x += int(w * scale)
    for (x0, w), h in zip(xs, headers):
        text(s, x0, top, w, Inches(0.4), [(h, 13, True, MUTED, SANS)])
    rule(s, Emu(int(top) + Inches(0.45)))
    for i, row in enumerate(rows):
        ry = Emu(int(top) + Inches(0.57) + Inches(0.46) * i)
        for j, ((x0, w), cell) in enumerate(zip(xs, row)):
            bold = isinstance(cell, tuple)
            val = cell[0] if bold else cell
            color = ACCENT if bold else INK
            font = MONO if j > 0 else SANS
            text(s, x0, ry, w, Inches(0.4), [(val, 15, bold, color, font)])
    if note:
        text(s, M, H - Inches(1.15), W - 2 * M, Inches(0.7),
             [(note, 14, False, MUTED, SANS)], spacing=1.15)
    return s


def figure_slide(prs, kicker, title, img, cap=None, height=Inches(4.3),
                 top=Inches(1.95)):
    s = blank(prs)
    header(s, kicker, title)
    picture(s, img, top=top, height=height)
    if cap:
        caption(s, cap)
    return s


def section_slide(prs, kicker, title, sub=None):
    s = blank(prs)
    text(s, M, Inches(2.6), W - 2 * M, Inches(0.4),
         [(kicker.upper(), 14, True, ACCENT, SANS)])
    text(s, M, Inches(3.0), W - 2 * M, Inches(1.2), [(title, 40, True, BLUE, SANS)])
    rule(s, Inches(4.3), w=Inches(3.0))
    if sub:
        text(s, M, Inches(4.6), Inches(9.8), Inches(1.2),
             [(sub, 18, False, MUTED, SANS)], spacing=1.2)
    return s




# ------------------------------------------------------------------ equations
# Every symbol used here is explained in words on the slide that shows it.

EQ_SCHEDULE = (r"$\beta_t:\ 10^{-4} \rightarrow 0.02,\qquad \alpha_t = 1-\beta_t,"
               r"\qquad \bar\alpha_t = \alpha_1\,\alpha_2 \cdots \alpha_t$")
EQ_FORWARD = (r"$x_t=\sqrt{\bar\alpha_t}\,x_0+\sqrt{1-\bar\alpha_t}\,\epsilon,"
              r"\qquad \epsilon\sim\mathcal{N}(0,\,1)$")
EQ_LOSS = (r"$\mathcal{L}=\mathbb{E}_{\,x_0,\,y,\,t,\,\epsilon}"
           r"\left[\ \|\,\epsilon-\epsilon_\theta(x_t,\,t,\,y)\,\|^2\ \right]$")
EQ_BLOCK = r"$h=\mathrm{conv}_2\left(\ \mathrm{conv}_1(x)\ +\ W\,e\ \right)$"
EQ_TIME = (r"$\mathrm{PE}(t)=\left[\ \sin(\omega_i t),\ \cos(\omega_i t)\ \right]_{i=0}^{23},"
           r"\qquad \omega_i=10000^{-i/24}$")
EQ_EMB = r"$e=\mathrm{MLP}\left(\mathrm{PE}(t)\right)+E[y]$"
EQ_DDPM = (r"$x_{t-1}=\frac{1}{\sqrt{\alpha_t}}\left(x_t-\frac{\beta_t}{\sqrt{1-\bar\alpha_t}}"
           r"\,\hat\epsilon\right)+\sqrt{\beta_t}\,z$")
EQ_GUIDE = (r"$\hat\epsilon=\epsilon_\theta(x_t,t,\emptyset)+w\,\left(\epsilon_\theta(x_t,t,y)"
            r"-\epsilon_\theta(x_t,t,\emptyset)\right)$")
EQ_X0 = r"$\hat{x}_0=\frac{x_t-\sqrt{1-\bar\alpha_t}\ \hat\epsilon}{\sqrt{\bar\alpha_t}}$"
EQ_JUMP = (r"$x_{t'}=\sqrt{\bar\alpha_{t'}}\ \hat{x}_0+\sqrt{1-\bar\alpha_{t'}}\ \hat\epsilon,"
           r"\qquad t' < t$")
EQ_X0_FROM_EPS = (r"$\mathrm{know}\ \epsilon:\qquad x_0=\frac{x_t-\sqrt{1-\bar\alpha_t}\ \epsilon}"
                  r"{\sqrt{\bar\alpha_t}}$")
EQ_EPS_FROM_X0 = (r"$\mathrm{know}\ x_0:\qquad \epsilon=\frac{x_t-\sqrt{\bar\alpha_t}\ x_0}"
                  r"{\sqrt{1-\bar\alpha_t}}$")
EQ_WEIGHT = (r"$\|\,\hat\epsilon-\epsilon\,\|^2\ =\ \frac{\bar\alpha_t}{1-\bar\alpha_t}"
             r"\ \|\,\hat{x}_0-x_0\,\|^2$")
EQ_QKV = r"$Q=h\,W_Q,\qquad K=h\,W_K,\qquad V=h\,W_V$"
EQ_ATTN = r"$\mathrm{Attn}(h)=\mathrm{softmax}\left(\frac{Q\,K^{T}}{\sqrt{d}}\right)V$"
ALL_EQS = [EQ_SCHEDULE, EQ_FORWARD, EQ_LOSS, EQ_BLOCK, EQ_TIME, EQ_EMB, EQ_DDPM,
           EQ_GUIDE, EQ_X0, EQ_JUMP,
           EQ_X0_FROM_EPS, EQ_EPS_FROM_X0, EQ_WEIGHT, EQ_QKV, EQ_ATTN]


# ------------------------------------------------------------------ data

def load_json(name, folder=RESULTS):
    path = os.path.join(folder, name)
    if not os.path.exists(path):
        return None
    with open(path) as f:
        return json.load(f)


def result(name, folder=RESULTS):
    path = os.path.join(folder, name)
    return path if os.path.exists(path) else None


def trained_base():
    """The base width the checkpoint was trained with (48 if not trained yet)."""
    if os.path.exists(CKPT):
        return torch.load(CKPT, map_location="cpu")["config"]["base"]
    return 48


BLOCK_NAMES = ["enc1", "enc2", "enc3", "bottleneck", "dec3", "dec2", "dec1", "head"]


def real_shapes(model):
    """{block name: (C, H, W)} from forward hooks on a real forward pass."""
    shapes, hooks = {}, []
    for n in BLOCK_NAMES:
        hooks.append(getattr(model, n).register_forward_hook(
            lambda _m, _i, out, n=n: shapes.__setitem__(n, tuple(out.shape[1:]))))
    zero = torch.zeros(1, dtype=torch.long)
    with torch.no_grad():
        model(torch.zeros(1, 1, IMAGE_SIZE, IMAGE_SIZE), zero, zero)
    for h in hooks:
        h.remove()
    return shapes


def shape_str(s):
    return " x ".join(str(v) for v in s)


# ------------------------------------------------------------------ drawing helpers

def show_img(ax, x, title=None, fs=11):
    ax.imshow(((x.squeeze() + 1) / 2).clamp(0, 1).numpy(), cmap="gray", vmin=0, vmax=1)
    ax.set_xticks([]); ax.set_yticks([])
    for sp in ax.spines.values():
        sp.set_visible(False)
    if title:
        ax.set_title(title, fontsize=fs, color="#1c222b")


def box(ax, x, y, w, h, label, color, fs=10, tc="white", alpha=0.92):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                fc=color, ec="none", alpha=alpha))
    ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=fs, color=tc)


def arrow(ax, start, end, color=C_GREY, lw=1.6, ls="-", rad=0.0):
    ax.annotate("", xy=end, xytext=start,
                arrowprops=dict(arrowstyle="-|>", color=color, lw=lw, linestyle=ls,
                                shrinkA=0, shrinkB=0, connectionstyle=f"arc3,rad={rad}"))


def canvas(w, h, xlim, ylim):
    fig = plt.figure(figsize=(w, h))
    fig.patch.set_facecolor(PAPER_HEX)
    ax = fig.add_axes([0, 0, 1, 1])              # data coordinates = the whole figure
    ax.set_facecolor(PAPER_HEX); ax.axis("off")
    ax.set_xlim(*xlim); ax.set_ylim(*ylim)
    return fig, ax


# ------------------------------------------------------------------ figures

def fig_forward(path, x0, sched):
    """A real training digit pushed to eight noise levels with the forward formula."""
    ts = [0, 50, 100, 200, 300, 500, 700, 999]
    noise = torch.randn(x0.shape, generator=torch.Generator().manual_seed(0))
    fig, axes = plt.subplots(1, len(ts), figsize=(12, 2.1))
    fig.patch.set_facecolor(PAPER_HEX)
    for ax, t in zip(axes, ts):
        xt = sched.add_noise(x0, torch.tensor([t]), noise)
        ab = sched.alpha_bars[t].item()
        show_img(ax, xt, f"t = {t}\n" + r"$\bar\alpha_t$" + f" = {ab:.2f}")
    return save(fig, path)


def fig_schedule(path, sched):
    ab = sched.alpha_bars
    fig, ax = plt.subplots(figsize=(8, 2.6))
    ax.plot(ab.sqrt().numpy(), color=C_COOL, lw=2.2, label=r"$\sqrt{\bar\alpha_t}$  (how much digit)")
    ax.plot((1 - ab).sqrt().numpy(), color=C_WARM, lw=2.2, label=r"$\sqrt{1-\bar\alpha_t}$  (how much noise)")
    style(ax, "noise level t", "weight")
    ax.legend(frameon=False, loc="center right")
    return save(fig, path)


def fig_io(path, x0, sched, n_params):
    """Input -> U-Net -> output, with a real noised digit and its real noise."""
    t = 300
    noise = torch.randn(x0.shape, generator=torch.Generator().manual_seed(0))
    xt = sched.add_noise(x0, torch.tensor([t]), noise)

    fig, ax = canvas(12, 3.8, (0, 12), (0, 3.8))
    for left, img, label in ((0.2, xt, "input  $x_t$: the noisy digit\n1 x 32 x 32"),
                             (8.6, noise, "target  $\\epsilon$: the noise we added\n(the output $\\hat\\epsilon$ must match it)")):
        sub = fig.add_axes([left / 12, 0.26, 2.4 / 12, 2.4 / 3.8])
        show_img(sub, img)
        ax.text(left + 1.2, 0.55, label, ha="center", va="top", fontsize=11, color="#1c222b")

    box(ax, 3.2, 2.35, 1.5, 0.6, f"t = {t}", C_COOL, fs=12)
    ax.text(3.95, 3.1, "noise level", ha="center", fontsize=10, color=C_GREY)
    box(ax, 3.2, 1.05, 1.5, 0.6, "y = 7", C_COOL, fs=12)
    ax.text(3.95, 0.75, "which digit", ha="center", va="top", fontsize=10, color=C_GREY)

    box(ax, 5.5, 0.9, 2.3, 2.2, f"U-Net\n{n_params:,}\nparameters", C_WARM, fs=13)
    arrow(ax, (2.65, 2.0), (5.45, 2.0))
    arrow(ax, (4.75, 2.65), (5.45, 2.35))
    arrow(ax, (4.75, 1.35), (5.45, 1.65))
    arrow(ax, (7.85, 2.0), (8.55, 2.0))
    return save(fig, path)


def image_at(fig, ax, x, y, size, img, xlim, ylim):
    """Place a square image, `size` data units wide, lower-left corner at (x, y)."""
    (x0, x1), (y0, y1) = xlim, ylim
    fw, fh = fig.get_size_inches()
    inches = size * fw / (x1 - x0)
    sub = fig.add_axes([(x - x0) / (x1 - x0), (y - y0) / (y1 - y0), inches / fw, inches / fh])
    show_img(sub, img)
    return sub


LEARNED = C_WARM        # training diagram: weights that backprop updates
FROZEN = C_COOL         # inference diagram: the same weights, now fixed
DATA = "#9aa3ad"        # data and fixed arithmetic: nothing to learn


def fig_training_flow(path, x0, sched, emb_dim):
    """One training step: where every input comes from, and what gets updated."""
    xlim, ylim = (0, 16), (-1.3, 7.0)
    fig, ax = canvas(14, 14 * 8.3 / 16, xlim, ylim)
    t = 300
    noise = torch.randn(x0.shape, generator=torch.Generator().manual_seed(0))
    xt = sched.add_noise(x0, torch.tensor([t]), noise)

    # inputs, left column
    image_at(fig, ax, 0.3, 5.0, 1.5, x0, xlim, ylim)
    ax.text(1.05, 4.75, "real digit $x_0$\n(from MNIST)", ha="center", va="top", fontsize=10)
    image_at(fig, ax, 0.3, 2.75, 1.5, noise, xlim, ylim)
    ax.text(1.05, 2.5, "random noise $\\epsilon$", ha="center", va="top", fontsize=10)
    box(ax, 0.2, 1.3, 1.7, 0.6, f"t = {t}\n(random)", DATA, fs=10)
    box(ax, 0.2, 0.05, 1.7, 0.6, "y = 7\n(its true label)", DATA, fs=10)

    # the forward formula: fixed arithmetic, nothing learned
    box(ax, 2.6, 4.2, 2.6, 1.3, "add noise\n$x_t=\\sqrt{\\bar\\alpha_t}\\,x_0+\\sqrt{1-\\bar\\alpha_t}\\,\\epsilon$",
        DATA, fs=10)
    arrow(ax, (1.85, 5.75), (2.6, 5.1))
    arrow(ax, (1.85, 3.5), (2.6, 4.6))
    arrow(ax, (1.9, 1.75), (3.0, 4.2))
    image_at(fig, ax, 5.8, 4.1, 1.5, xt, xlim, ylim)
    ax.text(6.55, 3.85, "noisy digit $x_t$", ha="center", va="top", fontsize=10)
    arrow(ax, (5.2, 4.85), (5.8, 4.85))

    # the two embeddings: learned
    box(ax, 2.6, 1.3, 2.6, 0.6, "sin/cos + MLP", LEARNED, fs=10)
    box(ax, 2.6, 0.0, 2.6, 0.7, "label table (11 x 192)\nrow 7", LEARNED, fs=10)
    arrow(ax, (1.9, 1.6), (2.6, 1.6))
    arrow(ax, (1.9, 0.35), (2.6, 0.35))
    ax.add_patch(plt.Circle((5.9, 0.95), 0.25, color=C_GREY))
    ax.text(5.9, 0.95, "+", ha="center", va="center", fontsize=15, color="white")
    arrow(ax, (5.2, 1.6), (5.72, 1.12))
    arrow(ax, (5.2, 0.35), (5.72, 0.78))
    box(ax, 6.6, 0.65, 1.6, 0.6, f"e\n{emb_dim} numbers", "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    arrow(ax, (6.15, 0.95), (6.6, 0.95))

    # the network
    box(ax, 8.8, 1.8, 2.0, 3.7, "U-Net", LEARNED, fs=14)
    arrow(ax, (7.3, 4.85), (8.8, 4.85))
    arrow(ax, (8.2, 0.95), (9.3, 1.8), rad=0.25)
    box(ax, 11.4, 3.2, 1.7, 0.9, "$\\hat\\epsilon$\npredicted noise", DATA, fs=10)
    arrow(ax, (10.8, 3.65), (11.4, 3.65))

    # the loss
    box(ax, 13.6, 5.2, 2.2, 0.9, "target: $\\epsilon$\n(the noise we added)", DATA, fs=10)
    box(ax, 13.6, 3.2, 2.2, 0.9, "loss\nMSE($\\hat\\epsilon$, $\\epsilon$)", "#3d4a5c", fs=11)
    arrow(ax, (14.7, 5.2), (14.7, 4.1))
    arrow(ax, (13.1, 3.65), (13.6, 3.65))

    # backprop
    red = "#b0301e"
    ax.plot([14.7, 14.7, 3.9], [3.2, -0.55, -0.55], color=red, lw=1.8, ls="--")
    arrow(ax, (3.9, -0.55), (3.9, 0.0), color=red, lw=1.8, ls="--")
    ax.plot([9.8, 9.8], [-0.55, 0.3], color=red, lw=1.8, ls="--")
    arrow(ax, (9.8, 0.3), (9.8, 1.8), color=red, lw=1.8, ls="--")
    ax.text(12.3, -0.95, "backprop: the loss updates everything orange -- the U-Net, the MLP "
            "AND the label table", ha="center", fontsize=10.5, color=red)
    ax.text(1.05, -0.3, "10% of the time y is\nreplaced by \"no label\"", ha="center", va="top",
            fontsize=8.5, color=C_GREY)
    ax.text(0.2, 6.85, "orange = learned (changed by training)     grey = data or fixed arithmetic",
            fontsize=10, color="#1c222b", va="top")
    return save(fig, path)


def fig_inference_flow(path, emb_dim):
    """Generating a 7: the same pieces, no loss, weights fixed, run in a loop."""
    xlim, ylim = (0, 16), (-1.3, 7.0)
    fig, ax = canvas(14, 14 * 8.3 / 16, xlim, ylim)
    start = torch.randn(1, 1, IMAGE_SIZE, IMAGE_SIZE, generator=torch.Generator().manual_seed(0))

    # inputs
    box(ax, 0.2, 5.9, 1.7, 0.55, "seed = 0", DATA, fs=10)
    image_at(fig, ax, 0.3, 3.8, 1.5, start, xlim, ylim)
    arrow(ax, (1.05, 5.9), (1.05, 5.33))
    ax.text(1.05, 3.55, "pure noise\n(start image)", ha="center", va="top", fontsize=10)
    box(ax, 0.2, 1.3, 1.7, 0.6, "t = 999, 998,\n..., 0", DATA, fs=10)
    box(ax, 0.2, 0.05, 1.7, 0.6, "y = 7\n(what YOU want)", DATA, fs=10)

    box(ax, 2.6, 1.3, 2.6, 0.6, "sin/cos + MLP", FROZEN, fs=10)
    box(ax, 2.6, 0.0, 2.6, 0.7, "label table (11 x 192)\nrow 7", FROZEN, fs=10)
    arrow(ax, (1.9, 1.6), (2.6, 1.6))
    arrow(ax, (1.9, 0.35), (2.6, 0.35))
    ax.add_patch(plt.Circle((5.9, 0.95), 0.25, color=C_GREY))
    ax.text(5.9, 0.95, "+", ha="center", va="center", fontsize=15, color="white")
    arrow(ax, (5.2, 1.6), (5.72, 1.12))
    arrow(ax, (5.2, 0.35), (5.72, 0.78))
    box(ax, 6.6, 0.65, 1.6, 0.6, f"e\n{emb_dim} numbers", "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    arrow(ax, (6.15, 0.95), (6.6, 0.95))

    # the loop
    ax.add_patch(FancyBboxPatch((5.3, 2.35), 8.3, 4.1, boxstyle="round,pad=0.05,rounding_size=0.2",
                                fc="none", ec=C_GREY, lw=1.4, ls="--"))
    ax.text(5.5, 2.5, "repeat 1000 times, t = 999 down to 0", fontsize=10.5, color=C_GREY, va="bottom")
    box(ax, 5.8, 3.9, 1.7, 1.1, "current\nimage $x_t$", DATA, fs=10)
    arrow(ax, (1.8, 4.55), (5.8, 4.45))
    box(ax, 8.8, 3.1, 2.0, 2.7, "U-Net", FROZEN, fs=14)
    arrow(ax, (7.5, 4.45), (8.8, 4.45))
    arrow(ax, (8.2, 0.95), (9.8, 3.1), rad=0.2)
    box(ax, 11.3, 3.75, 2.0, 1.4, "remove part of $\\hat\\epsilon$\n(+ a little\nfresh noise)", DATA, fs=9.5)
    arrow(ax, (10.8, 4.45), (11.3, 4.45))
    ax.text(11.05, 4.6, "$\\hat\\epsilon$", ha="center", fontsize=11)
    arrow(ax, (12.3, 5.15), (6.65, 5.0), rad=0.22)
    ax.text(9.5, 6.08, "the result becomes the next, slightly cleaner $x_t$", ha="center",
            fontsize=10, color="#1c222b")

    box(ax, 14.1, 3.9, 1.8, 1.1, "a new 7\n1 x 32 x 32", "#3d4a5c", fs=11)
    arrow(ax, (13.3, 4.1), (14.1, 4.3))
    ax.text(15.0, 3.65, "after t = 0", ha="center", va="top", fontsize=9, color=C_GREY)

    ax.text(10.5, -0.35, "no target, no loss, no backprop:\nevery weight, including the label "
            "table, stays fixed", ha="center", va="top", fontsize=10.5, color=FROZEN)
    ax.text(0.2, 6.85, "blue = learned during training, now fixed", fontsize=10,
            color="#1c222b", va="top")
    return save(fig, path)


def fig_embedding(path, emb_dim):
    """t and y -> e: what the sin/cos + MLP block and the label table do."""
    xlim, ylim = (0, 16), (-0.45, 4.35)
    fig, ax = canvas(14, 14 * 4.8 / 16, xlim, ylim)
    t = 300

    # ---- t: fixed sin/cos formula, then a learned MLP
    ty = 3.25
    box(ax, 0.2, ty, 1.4, 0.6, f"t = {t}", DATA, fs=11)
    box(ax, 2.1, ty, 1.5, 0.6, "sin / cos\n(fixed formula)", DATA, fs=9.5)
    arrow(ax, (1.6, ty + 0.3), (2.1, ty + 0.3))
    pe = timestep_embedding(torch.tensor([t]), 48)
    strip = fig.add_axes([4.0 / 16, (ty + 0.05 - ylim[0]) / (ylim[1] - ylim[0]), 3.3 / 16,
                          0.5 / (ylim[1] - ylim[0])])
    strip.imshow(pe.numpy(), aspect="auto", cmap="RdBu_r", vmin=-1, vmax=1)
    strip.set_xticks([]); strip.set_yticks([])
    arrow(ax, (3.6, ty + 0.3), (4.0, ty + 0.3))
    ax.text(5.65, ty + 0.68, f"PE({t}): 48 numbers between -1 and 1", ha="center", fontsize=9.5)
    ax.text(5.65, ty - 0.12, "fast waves tell t = 300 from 301,\nslow waves tell 300 from 800",
            ha="center", va="top", fontsize=8.5, color=C_GREY)
    box(ax, 7.8, ty - 0.1, 2.8, 0.8, f"MLP (learned)\nLinear 48 -> {emb_dim}, ReLU,\n"
        f"Linear {emb_dim} -> {emb_dim}", LEARNED, fs=9.5)
    arrow(ax, (7.3, ty + 0.3), (7.8, ty + 0.3))
    box(ax, 11.2, ty, 1.6, 0.6, f"{emb_dim} numbers", "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    arrow(ax, (10.6, ty + 0.3), (11.2, ty + 0.3))

    # ---- y: a learned table, one row per label
    top, rh, x0, x1 = 2.25, 0.19, 2.9, 10.6
    names = [str(d) for d in range(10)] + ["no label"]
    for i, name in enumerate(names):
        y = top - (i + 1) * rh
        hit = i == 7
        ax.add_patch(plt.Rectangle((x0, y + 0.015), x1 - x0, rh - 0.03,
                                   color=LEARNED if hit else "#e6ded5"))
        ax.text(x0 - 0.1, y + rh / 2, name, ha="right", va="center", fontsize=8.5,
                color=LEARNED if hit else C_GREY, fontweight="bold" if hit else "normal")
    ax.text((x0 + x1) / 2, -0.12, f"label table (learned): 11 rows x {emb_dim} numbers  "
            "=  nn.Embedding(11, " + f"{emb_dim})", ha="center", va="top", fontsize=9.5)
    row7 = top - 8 * rh + rh / 2
    box(ax, 0.2, row7 - 0.3, 1.4, 0.6, "y = 7", DATA, fs=11)
    arrow(ax, (1.6, row7), (2.05, row7))
    box(ax, 11.2, row7 - 0.3, 1.6, 0.6, f"row 7:\n{emb_dim} numbers", "#e8d6c8", fs=10,
        tc="#1c222b", alpha=1.0)
    arrow(ax, (10.6, row7), (11.2, row7))

    # ---- add them: e
    cx, cy = 13.55, (ty + 0.3 + row7) / 2
    ax.add_patch(plt.Circle((cx, cy), 0.25, color=C_GREY))
    ax.text(cx, cy, "+", ha="center", va="center", fontsize=15, color="white")
    arrow(ax, (12.8, ty + 0.3), (cx - 0.12, cy + 0.22))
    arrow(ax, (12.8, row7), (cx - 0.12, cy - 0.22))
    box(ax, 14.2, cy - 0.35, 1.7, 0.7, f"e\n{emb_dim} numbers", "#e8d6c8", fs=11, tc="#1c222b", alpha=1.0)
    arrow(ax, (cx + 0.25, cy), (14.2, cy))
    ax.text(15.05, cy - 0.5, "goes into every\nblock (below)", ha="center", va="top", fontsize=9,
            color=C_GREY)
    return save(fig, path)


def fig_loss_weight(path, x0, sched):
    """How much one unit of digit error costs in the noise loss, at every t,
    with the real noisy digit at five of those t shown above the curve."""
    from matplotlib.transforms import blended_transform_factory
    ab = sched.alpha_bars
    weight = (ab / (1 - ab)).numpy()
    noise = torch.randn(x0.shape, generator=torch.Generator().manual_seed(0))
    fig, ax = plt.subplots(figsize=(11, 3.9))
    fig.subplots_adjust(top=0.62)
    ax.plot(weight, color=C_WARM, lw=2.4)
    ax.set_yscale("log")
    style(ax, "noise level t", "weight  " + r"$\bar\alpha_t\,/\,(1-\bar\alpha_t)$")
    above = blended_transform_factory(ax.transData, ax.transAxes)
    for t in (0, 100, 300, 600, 999):
        ax.plot(t, weight[t], "o", color=C_WARM, ms=6)
        w = weight[t]
        label = (f"{w:,.0f}" if w >= 100 else f"{w:.1f}" if w >= 1
                 else f"{w:.2g}" if w >= 1e-3 else f"{w:.5f}")
        ax.annotate(label, (t, w), xytext=(8, 6), textcoords="offset points",
                    fontsize=10, color="#1c222b")
        thumb = ax.inset_axes([t - 40, 1.06, 80, 0.5], transform=above)
        show_img(thumb, sched.add_noise(x0, torch.tensor([t]), noise), f"t = {t}", fs=9.5)
    ax.set_xlim(-60, 1060)
    return save(fig, path)


def fig_unet(path, shapes, n_params, base):
    """The U-Net drawn as a U, every shape taken from a real forward pass."""
    fig, ax = canvas(13, 5.4, (-0.3, 15.1), (-1.75, 4.0))
    w, h = 1.3, 0.66
    pos = {"enc1": (1.5, 3), "enc2": (3.0, 2), "enc3": (4.5, 1), "bottleneck": (6.0, 0),
           "dec3": (7.5, 1), "dec2": (9.0, 2), "dec1": (10.5, 3), "head": (12.0, 3)}
    colour = {n: C_COOL for n in ("enc1", "enc2", "enc3")}
    colour.update(bottleneck="#3d4a5c", dec3=C_GREY, dec2=C_GREY, dec1=C_GREY, head=C_WARM)

    box(ax, 0.0, 3, w, h, "input $x_t$\n1 x 32 x 32", "#9aa3ad", fs=9.5)
    box(ax, 13.5, 3, w, h, "output $\\hat\\epsilon$\n1 x 32 x 32", C_WARM, fs=9.5)
    for n, (x, y) in pos.items():
        box(ax, x, y, w, h, f"{n}\n{shape_str(shapes[n])}", colour[n], fs=9.5)

    def mid(n, side):
        x, y = pos[n]
        return {"l": (x, y + h / 2), "r": (x + w, y + h / 2),
                "b": (x + w / 2, y), "t": (x + w / 2, y + h)}[side]

    arrow(ax, (w, 3 + h / 2), mid("enc1", "l"))
    arrow(ax, mid("head", "r"), (13.5, 3 + h / 2))
    arrow(ax, mid("dec1", "r"), mid("head", "l"))
    for a, b in (("enc1", "enc2"), ("enc2", "enc3"), ("enc3", "bottleneck")):
        arrow(ax, mid(a, "b"), mid(b, "l"), color=C_COOL)
    for a, b in (("bottleneck", "dec3"), ("dec3", "dec2"), ("dec2", "dec1")):
        arrow(ax, mid(a, "r"), mid(b, "b"), color=C_GREY)
    ax.text(3.2, 1.35, "pool\n(halve)", ha="center", fontsize=9, color=C_COOL)
    ax.text(9.4, 1.35, "up\n(double)", ha="center", fontsize=9, color=C_GREY)
    for a, b in (("enc1", "dec1"), ("enc2", "dec2"), ("enc3", "dec3")):
        arrow(ax, mid(a, "r"), mid(b, "l"), color=C_WARM, lw=1.8, ls="--")
    ax.text(6.65, 3 + h / 2 + 0.1, "skip: copy across (torch.cat)", ha="center",
            fontsize=9.5, color=C_WARM)

    # the t-and-y vector, fed into every block
    box(ax, 1.5, -1.6, 11.8, 0.5, f"e: {base * 4} numbers made from t and y  ->  added inside every block",
        "#e8d6c8", fs=10.5, tc="#1c222b", alpha=1.0)
    for n, (x, y) in pos.items():
        if n != "head":
            ax.plot([x + w / 2, x + w / 2], [-1.1, y], color=C_WARM, lw=0.9, ls=":", alpha=0.8)
    ax.set_title(f"week 5 U-Net, base {base}: {n_params:,} parameters", fontsize=12, color="#1c222b")
    return save(fig, path)


def fig_block(path, shapes, base):
    """Inside one Block (enc2 as the example), with its real channel counts."""
    c_in, c_out = shapes["enc1"][0], shapes["enc2"][0]
    hw = f"{shapes['enc2'][1]} x {shapes['enc2'][2]}"
    fig, ax = canvas(13, 3.2, (0, 13.4), (-0.2, 3.1))
    y, h = 1.9, 0.75
    steps = [(0.0, 1.5, f"input\n{c_in} x {hw}", "#9aa3ad"),
             (2.0, 1.9, f"conv 3x3\nGroupNorm, ReLU", C_COOL),
             (5.5, 0.0, "+", C_WARM),
             (7.2, 1.9, f"conv 3x3\nGroupNorm, ReLU", C_COOL),
             (9.8, 1.9, f"output\n{c_out} x {hw}", "#9aa3ad")]
    for x, w, label, col in steps:
        if label == "+":
            ax.add_patch(plt.Circle((x + 0.35, y + h / 2), 0.35, color=C_WARM))
            ax.text(x + 0.35, y + h / 2, "+", ha="center", va="center", fontsize=20, color="white")
        else:
            box(ax, x, y, w, h, label, col, fs=10.5)
    arrow(ax, (1.5, y + h / 2), (2.0, y + h / 2))
    arrow(ax, (3.9, y + h / 2), (5.5, y + h / 2))
    ax.text(4.7, y + h / 2 + 0.12, f"{c_out} x {hw}", ha="center", fontsize=9.5, color=C_GREY)
    arrow(ax, (6.2, y + h / 2), (7.2, y + h / 2))
    arrow(ax, (9.1, y + h / 2), (9.8, y + h / 2))

    box(ax, 0.0, 0.1, 1.5, 0.7, f"e\n{base * 4} numbers", "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    box(ax, 2.2, 0.1, 2.4, 0.7, f"Linear  {base * 4} -> {c_out}", C_WARM, fs=10.5)
    arrow(ax, (1.5, 0.45), (2.2, 0.45))
    arrow(ax, (4.6, 0.45), (5.85, y - 0.02), color=C_WARM, rad=0.25)
    ax.text(6.3, 0.5, f"{c_out} numbers, one per channel,\nsame value at all {hw} pixels",
            fontsize=9.5, color=C_WARM)
    return save(fig, path)


def fig_time_embedding(path):
    t = torch.arange(1000)
    pe = timestep_embedding(t, 48)
    fig, ax = plt.subplots(figsize=(10, 2.6))
    im = ax.imshow(pe.T.numpy(), aspect="auto", cmap="RdBu_r", vmin=-1, vmax=1,
                   interpolation="nearest")
    style(ax, "noise level t", "number (0..47)")
    ax.grid(False)
    fig.colorbar(im, ax=ax, pad=0.01)
    return save(fig, path)


def fig_receptive_field(path, model, x0):
    """Which input pixels can change ONE feature at the centre of each level?

    Measured with gradients: d(feature) / d(input) is non-zero exactly at the
    pixels that feature can see. GroupNorm is switched off for this, because
    it mixes in an image-wide average and would make every pixel count.
    """
    m = copy.deepcopy(model)
    for blk in m.modules():
        if isinstance(blk, Block):
            blk.conv1[1] = nn.Identity()
            blk.conv2[1] = nn.Identity()
    feats = {}
    for n in ("enc1", "enc2", "enc3", "bottleneck"):
        getattr(m, n).register_forward_hook(lambda _m, _i, out, n=n: feats.__setitem__(n, out))

    x = torch.randn(8, 1, IMAGE_SIZE, IMAGE_SIZE, requires_grad=True)
    zero = torch.zeros(8, dtype=torch.long)
    m(x, zero, zero)
    fig, axes = plt.subplots(1, 4, figsize=(12, 3.3))
    fig.patch.set_facecolor(PAPER_HEX)
    for ax, n in zip(axes, ("enc1", "enc2", "enc3", "bottleneck")):
        f = feats[n]
        grad, = torch.autograd.grad(f[:, :, f.shape[2] // 2, f.shape[3] // 2].sum(), x,
                                    retain_graph=True)
        seen = grad.abs().sum(0)[0] > 0
        rows, cols = seen.any(1).nonzero(), seen.any(0).nonzero()
        side_h, side_w = int(rows.max() - rows.min() + 1), int(cols.max() - cols.min() + 1)
        show_img(ax, x0)
        overlay = torch.zeros(IMAGE_SIZE, IMAGE_SIZE, 4)
        overlay[seen] = torch.tensor([0.76, 0.33, 0.12, 0.55])
        ax.imshow(overlay.numpy())
        sees = "the whole image" if seen.all() else f"{side_h} x {side_w} pixels"
        ax.set_title(f"{n}  ({f.shape[2]} x {f.shape[3]})\nsees {sees}",
                     fontsize=11, color="#1c222b")
    return save(fig, path)


def fig_loss(path, log):
    fig, ax = plt.subplots(figsize=(5.4, 3.8))
    ax.plot(log["epoch"], log["loss"], marker="o", color=C_WARM, lw=2)
    style(ax, "epoch", "training loss (MSE)")
    ax.set_xticks([e for e in log["epoch"] if e == 1 or e % 5 == 0])
    return save(fig, path)


def fig_images(path, items, height=4.0):
    """Saved result PNGs side by side, each with a title."""
    fig, axes = plt.subplots(1, len(items), figsize=(height * len(items), height + 0.5))
    fig.patch.set_facecolor(PAPER_HEX)
    for ax, (p, title) in zip(axes if len(items) > 1 else [axes], items):
        ax.imshow(plt.imread(p), cmap="gray")
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values():
            sp.set_visible(False)
        ax.set_title(title, fontsize=12, color="#1c222b")
    return save(fig, path)


# ------------------------------------------------------------------ part 2 figures

def load_dit():
    """Part 2's checkpoint if it exists, else an untrained DiT."""
    if os.path.exists(CKPT2):
        ckpt = torch.load(CKPT2, map_location="cpu")
        c = ckpt["config"]
        model = DiT(dim=c["dim"], depth=c["depth"], heads=c["heads"], patch=c["patch"])
        model.load_state_dict(ckpt["model"])
        return model.eval()
    return DiT().eval()


def dit_shapes(model):
    """{stage: shape} from forward hooks on a real forward pass of the DiT."""
    shapes, hooks = {}, []
    stages = {"patchify": model.patchify, "block1": model.blocks[0],
              "block_last": model.blocks[-1], "head": model.head}
    for n, mod in stages.items():
        hooks.append(mod.register_forward_hook(
            lambda _m, _i, out, n=n: shapes.__setitem__(n, tuple(out.shape[1:]))))
    zero = torch.zeros(1, dtype=torch.long)
    with torch.no_grad():
        shapes["output"] = tuple(model(torch.zeros(1, 1, IMAGE_SIZE, IMAGE_SIZE), zero, zero).shape[1:])
    for h in hooks:
        h.remove()
    return shapes


def gflops(model):
    """Billions of multiply/add operations for ONE image, forward pass, counted by PyTorch."""
    from torch.utils.flop_counter import FlopCounterMode
    one = torch.zeros(1, dtype=torch.long)
    with FlopCounterMode(display=False) as fc, torch.no_grad():
        model(torch.zeros(1, 1, IMAGE_SIZE, IMAGE_SIZE), one, one)
    return fc.get_total_flops() / 1e9


def example_patch(x0, p):
    """(row, col) of the patch whose pixels vary most: the edge of a stroke,
    so its 4 numbers are not all the same."""
    patches = x0[0, 0].unfold(0, p, p).unfold(1, p, p)        # g x g x p x p
    g = patches.shape[0]
    i = int(patches.reshape(g, g, -1).std(-1).flatten().argmax())
    return i // g, i % g


def fig_patches(path, model, x0):
    """A real 7 cut into 2 x 2 patches; one patch becomes one token."""
    dim, p = model.dim, model.patch
    g = IMAGE_SIZE // p
    pr, pc = example_patch(x0, p)
    xlim, ylim = (0, 16), (-0.2, 5.2)
    fig, ax = canvas(14, 14 * 5.4 / 16, xlim, ylim)

    sub = image_at(fig, ax, 0.2, 0.6, 4.2, x0, xlim, ylim)
    for k in range(1, g):
        sub.axhline(k * p - 0.5, color=C_WARM, lw=0.35, alpha=0.7)
        sub.axvline(k * p - 0.5, color=C_WARM, lw=0.35, alpha=0.7)
    r0, c0 = pr * p, pc * p
    sub.add_patch(plt.Rectangle((c0 - 0.5, r0 - 0.5), p, p, fill=False, ec=C_WARM, lw=2.2))
    ax.text(2.3, 0.35, f"the 32 x 32 digit, cut into {g} x {g} = {g * g} patches\nof {p} x {p} pixels",
            ha="center", va="top", fontsize=10)

    # the one patch, blown up, with its 4 real pixel values
    patch = x0[0, 0, r0:r0 + p, c0:c0 + p]
    z = fig.add_axes([5.2 / 16, (2.0 - ylim[0]) / (ylim[1] - ylim[0]), 1.6 / 16,
                      1.6 * 14 / 16 / (14 * 5.4 / 16) ])
    z.imshow(((patch + 1) / 2).numpy(), cmap="gray", vmin=0, vmax=1)
    for (i, j), v in zip([(0, 0), (0, 1), (1, 0), (1, 1)], patch.flatten().tolist()):
        z.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=9,
               color="#1c222b" if v > 0 else "white")
    z.set_xticks([]); z.set_yticks([])
    for sp in z.spines.values():
        sp.set_color(C_WARM); sp.set_linewidth(2.2)
    arrow(ax, (4.45, 3.2), (5.15, 3.0), color=C_WARM)
    ax.text(6.0, 1.75, f"one patch:\n{p * p} numbers", ha="center", va="top", fontsize=10)

    box(ax, 7.3, 2.55, 2.3, 0.9, f"Linear {p * p} -> {dim}\n(learned)", LEARNED, fs=10)
    arrow(ax, (6.85, 3.0), (7.3, 3.0))
    ax.add_patch(plt.Circle((10.15, 3.0), 0.25, color=C_GREY))
    ax.text(10.15, 3.0, "+", ha="center", va="center", fontsize=15, color="white")
    arrow(ax, (9.6, 3.0), (9.9, 3.0))
    box(ax, 9.0, 0.75, 2.3, 0.9, f"position table\nrow {pr * g + pc} of {g * g} (learned)",
        LEARNED, fs=9.5)
    arrow(ax, (10.15, 1.65), (10.15, 2.75))
    box(ax, 10.8, 2.6, 1.6, 0.8, f"one token\n{dim} numbers", "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    arrow(ax, (10.4, 3.0), (10.8, 3.0))

    # all tokens: the real 256 x dim matrix the first block receives
    with torch.no_grad():
        tokens = model.patchify(x0).flatten(2).transpose(1, 2)[0] + model.pos[0]
    m = fig.add_axes([13.0 / 16, (0.6 - ylim[0]) / (ylim[1] - ylim[0]), 2.6 / 16,
                      4.2 / (ylim[1] - ylim[0])])
    lim = float(tokens.abs().quantile(0.98))
    m.imshow(tokens.numpy(), aspect="auto", cmap="RdBu_r", vmin=-lim, vmax=lim, interpolation="nearest")
    m.set_xticks([]); m.set_yticks([])
    arrow(ax, (12.4, 3.0), (12.95, 3.0))
    ax.text(14.3, 0.35, f"all {g * g} tokens x {dim} numbers:\nwhat the first block receives",
            ha="center", va="top", fontsize=10)
    ax.text(14.3, 5.05, "one row = one patch", ha="center", va="top", fontsize=9, color=C_GREY)
    return save(fig, path)


def draw_matrix(ax, x, top, m, cell, name, cmap="RdBu_r", lim=None, row=None, fs=9):
    """A small matrix as coloured cells with its numbers; `row` gets an orange box."""
    n, k = m.shape
    lim = lim or float(m.abs().max())
    colors = matplotlib.colormaps[cmap]
    for i in range(n):
        for j in range(k):
            v = float(m[i, j])
            c = colors(0.5 + 0.5 * v / lim) if cmap == "RdBu_r" else colors(0.15 + 0.7 * v / lim)
            ax.add_patch(plt.Rectangle((x + j * cell, top - (i + 1) * cell), cell, cell,
                                       fc=c, ec="white", lw=1.2, alpha=0.75))
            ax.text(x + (j + 0.5) * cell, top - (i + 0.5) * cell, f"{v:.2f}",
                    ha="center", va="center", fontsize=fs, color="#1c222b")
    if row is not None:
        ax.add_patch(plt.Rectangle((x, top - (row + 1) * cell), k * cell, cell,
                                   fill=False, ec=C_WARM, lw=2.4))
    ax.text(x + k * cell / 2, top + 0.12, name, ha="center", va="bottom", fontsize=10.5,
            color="#1c222b")
    return x + k * cell, top - n * cell / 2                  # right edge, vertical middle


def toy_attention(n=4, d=3, seed=0):
    """Real numbers for the toy example: random tokens and random W_Q, W_K, W_V."""
    g = torch.Generator().manual_seed(seed)
    h = torch.randn(n, d, generator=g)
    wq, wk, wv = (torch.randn(d, d, generator=g) for _ in range(3))
    q, k, v = h @ wq, h @ wk, h @ wv
    scores = q @ k.T / d ** 0.5
    a = torch.softmax(scores, dim=-1)
    return h, q, k, v, scores, a, a @ v


def fig_attention_toy(path, row=2):
    """Attention on 4 tokens of 3 numbers, every number computed; one token followed.
    Drawn in slide inches (1 unit = 1 inch) so the numbers stay readable."""
    h, q, k, v, scores, a, out = toy_attention()
    n, d = h.shape
    fig, ax = canvas(12.6, 4.75, (0, 12.6), (0.3, 5.05))
    c = 0.42
    lim = float(torch.cat([h, q, k, v, out]).abs().max())

    hr, hm = draw_matrix(ax, 0.1, 3.5, h, c, f"tokens h  ({n} x {d})", lim=lim, row=row)
    ax.text(0.1 + d * c / 2, 3.5 - n * c - 0.08, "one row = one token", ha="center", va="top",
            fontsize=8.5, color=C_GREY)
    qr, qm = draw_matrix(ax, 2.1, 4.65, q, c, "Q = h W_Q", lim=lim, row=row)
    kr, km = draw_matrix(ax, 2.1, 2.4, k, c, "K = h W_K", lim=lim)
    arrow(ax, (hr + 0.08, hm + 0.1), (2.03, qm), lw=1.2)
    arrow(ax, (hr + 0.08, hm - 0.1), (2.03, km), lw=1.2)

    sr, sm = draw_matrix(ax, 4.3, 3.5, scores, c, "scores = Q K$^T$ / $\\sqrt{3}$", row=row)
    arrow(ax, (qr + 0.08, qm), (4.23, sm + 0.25), lw=1.2)
    arrow(ax, (kr + 0.08, km), (4.23, sm - 0.25), lw=1.2)
    ax.text(4.3 + n * c / 2, 3.5 - n * c - 0.08, f"row {row + 1}: token {row + 1}'s query\n"
            "dotted with every key", ha="center", va="top", fontsize=8.5, color=C_GREY)

    ar, am = draw_matrix(ax, 6.9, 4.65, a, c, "weights = softmax(each row)", cmap="Oranges",
                         lim=1.0, row=row)
    arrow(ax, (sr + 0.08, sm), (6.83, am - 0.2), lw=1.2)
    for i in range(n):
        ax.text(ar + 0.08, 4.65 - (i + 0.5) * c, f"sum {float(a[i].sum()):.2f}", va="center",
                fontsize=8, color=C_GREY)

    vr, vm = draw_matrix(ax, 7.05, 2.1, v, c, "V = h W_V", lim=lim)
    ax.plot([0.1 + d * c / 2, 0.1 + d * c / 2, 6.6], [3.5 - n * c - 0.4, 0.45, 0.45],
            color=C_GREY, lw=1.2)
    arrow(ax, (6.6, 0.45), (6.98, vm - 0.3), lw=1.2)
    ax.text(3.6, 0.52, "W_Q, W_K, W_V: learned 3 x 3 matrices", ha="center", va="bottom",
            fontsize=9, color=LEARNED)

    orr, om = draw_matrix(ax, 10.4, 3.4, out, c, "output = weights V", lim=lim, row=row)
    arrow(ax, (ar + 0.8, am), (10.33, om + 0.25), lw=1.2)
    arrow(ax, (vr + 0.08, vm), (10.33, om - 0.25), lw=1.2)
    ax.text(10.4 + d * c / 2, 3.4 - n * c - 0.08, "same shape as h:\none row per token",
            ha="center", va="top", fontsize=8.5, color=C_GREY)
    return save(fig, path)


def blank_patches(x0, p, far_from):
    """Token numbers of all-black patches (every pixel -1), and the one furthest
    from token `far_from` that sits inside the digit's bounding box."""
    g = IMAGE_SIZE // p
    patches = x0[0, 0].unfold(0, p, p).unfold(1, p, p).reshape(g * g, -1)
    blank = (patches == -1).all(1)
    ink = (~blank).nonzero().flatten()
    r0, r1 = int(ink.min()) // g, int(ink.max()) // g
    c0, c1 = int((ink % g).min()), int((ink % g).max())
    inside = [i for i in blank.nonzero().flatten().tolist()
              if r0 < i // g < r1 and c0 < i % g < c1]
    fr, fc = divmod(far_from, g)
    return max(inside, key=lambda i: abs(i // g - fr) + abs(i % g - fc))


def position_stats(model):
    """Cosine similarity of the learned position rows: matrix, neighbour mean, far mean."""
    g = IMAGE_SIZE // model.patch
    rows = torch.nn.functional.normalize(model.pos[0].detach(), dim=1)
    sim = rows @ rows.T
    rc = torch.stack(torch.meshgrid(torch.arange(g), torch.arange(g), indexing="ij"), -1).reshape(-1, 2)
    steps = (rc[:, None, :] - rc[None, :, :]).abs().sum(-1)   # grid steps between two patches
    return sim, float(sim[steps == 1].mean()), float(sim[steps >= 6].mean())


def fig_position(path, model, x0, a=17):
    """Two all-black patches give identical tokens; the position table tells them apart."""
    p, g = model.patch, IMAGE_SIZE // model.patch
    b = blank_patches(x0, p, a)
    xlim, ylim = (0, 16), (-0.3, 5.4)
    fig, ax = canvas(14, 14 * 5.7 / 16, xlim, ylim)

    sub = image_at(fig, ax, 0.2, 0.75, 3.9, x0, xlim, ylim)
    for k in range(1, g):
        sub.axhline(k * p - 0.5, color=C_GRID, lw=0.3, alpha=0.6)
        sub.axvline(k * p - 0.5, color=C_GRID, lw=0.3, alpha=0.6)
    for tok, name, col in ((a, "A", C_COOL), (b, "B", C_WARM)):
        r, c = divmod(tok, g)
        sub.add_patch(plt.Rectangle((c * p - 0.5, r * p - 0.5), p, p, fill=False, ec=col, lw=2.4))
        sub.text(c * p + 0.5, r * p - 1.2, name, ha="center", va="bottom", fontsize=11,
                 color=col, fontweight="bold")
    ax.text(2.15, 0.55, "A and B: all black, every pixel -1", ha="center", va="top", fontsize=10)

    with torch.no_grad():
        tokens = model.patchify(x0).flatten(2).transpose(1, 2)[0]
    pos = model.pos[0].detach()
    k, cell = 4, 0.37
    ch = (pos[a] - pos[b]).abs().topk(k).indices.sort().values   # where the two rows differ most
    tokens, pos = tokens[:, ch], pos[:, ch]
    lim = float(torch.cat([tokens[[a, b]], pos[[a, b]], (tokens + pos)[[a, b]]]).abs().max())
    for tok, name, col, yc in ((a, "A", C_COOL, 3.75), (b, "B", C_WARM, 1.45)):
        r, c = divmod(tok, g)
        ax.text(4.45, yc + 0.32, f"{name}: token {tok}", ha="left", va="bottom", fontsize=10,
                color=col, fontweight="bold")
        ax.text(4.45, yc + 0.25, f"row {r}, col {c}", ha="left", va="top", fontsize=8.5, color=C_GREY)
        ax.text(4.45, yc - 0.3, "-1 -1 -1 -1", ha="left", va="center", fontsize=9, family="monospace")
        box(ax, 5.75, yc - 0.6, 1.0, 0.6, "Linear\n4 -> 176", LEARNED, fs=8.5)
        arrow(ax, (5.5, yc - 0.3), (5.75, yc - 0.3))
        arrow(ax, (6.75, yc - 0.3), (6.95, yc - 0.3))
        top = yc - 0.3 + cell / 2
        same = "patch numbers\n(the same for A and B)" if name == "A" else ""
        right, _ = draw_matrix(ax, 6.95, top, tokens[tok:tok + 1], cell, same, lim=lim, fs=7.5)
        ax.text(right + 0.25, yc - 0.3, "+", ha="center", va="center", fontsize=15, color=C_GREY)
        right, _ = draw_matrix(ax, right + 0.5, top, pos[tok:tok + 1], cell,
                               f"position row {tok}", lim=lim, fs=7.5)
        ax.text(right + 0.25, yc - 0.3, "=", ha="center", va="center", fontsize=15, color=C_GREY)
        draw_matrix(ax, right + 0.5, top, (tokens + pos)[tok:tok + 1], cell,
                    "what block 1 gets", lim=lim, fs=7.5)
    ax.text(9.2, 0.55, f"{k} of the {model.dim} numbers shown: numbers "
            + ", ".join(str(int(c)) for c in ch) + f",\nwhere rows {a} and {b} of the table differ most",
            ha="center", va="top", fontsize=9, color=C_GREY)

    sim, _, _ = position_stats(model)
    hm = fig.add_axes([12.75 / 16, (0.75 - ylim[0]) / (ylim[1] - ylim[0]), 3.1 / 16,
                       3.1 / (ylim[1] - ylim[0])])
    hm.imshow(sim[b].reshape(g, g).numpy(), cmap="Oranges", vmin=0, vmax=1)
    rb, cb = divmod(b, g)
    hm.add_patch(plt.Rectangle((cb - 0.5, rb - 0.5), 1, 1, fill=False, ec=C_COOL, lw=1.8))
    hm.set_xticks([]); hm.set_yticks([])
    ax.text(14.3, 4.15, f"after training: row {b} (B)\ncompared with every row",
            ha="center", va="bottom", fontsize=10)
    ax.text(14.3, 0.55, "dark = similar (1 = same direction)\nlight = unrelated (0)",
            ha="center", va="top", fontsize=9, color=C_GREY)
    return save(fig, path), a, b


def fig_plain_block(path, n_tok, dim, depth):
    """One standard transformer block: LN -> attention -> +, LN -> MLP -> +."""
    fig, ax = canvas(13, 2.6, (0, 16), (-0.6, 2.7))
    y, h = 0.6, 0.95
    stages = [(0.0, 1.5, f"tokens\n{n_tok} x {dim}", "#9aa3ad"),
              (2.0, 1.1, "LN", C_COOL),
              (3.6, 1.9, "attention", C_COOL),
              (6.6, 1.1, "LN", C_COOL),
              (8.2, 2.6, f"MLP, each token\n{dim} -> {4 * dim} -> {dim}", C_COOL),
              (14.4, 1.55, f"out\n{n_tok} x {dim}", "#9aa3ad")]
    for x, w, label, col in stages:
        box(ax, x, y, w, h, label, col, fs=10.5)
    mid = y + h / 2
    for cx, start in ((6.0, 1.5), (11.3, 6.1)):          # (+ circle, where its residual starts)
        ax.add_patch(plt.Circle((cx, mid), 0.28, color=C_GREY))
        ax.text(cx, mid, "+", ha="center", va="center", fontsize=17, color="white")
        ax.plot([start, start, cx], [mid, y + h + 0.45, y + h + 0.45], color=C_GREY, lw=1.3)
        arrow(ax, (cx, y + h + 0.45), (cx, mid + 0.28), lw=1.3)
    for (x, w, *_), (nx, *_) in zip(stages[:3], stages[1:3]):
        arrow(ax, (x + w, mid), (nx, mid))
    arrow(ax, (5.5, mid), (5.72, mid)); arrow(ax, (6.28, mid), (6.6, mid))
    arrow(ax, (7.7, mid), (8.2, mid)); arrow(ax, (10.8, mid), (11.02, mid))
    arrow(ax, (11.58, mid), (14.4, mid))
    ax.text(12.9, mid + 0.12, f"repeat: {depth} blocks", ha="center", va="bottom", fontsize=10, color=C_GREY)
    ax.text(4.55, y - 0.15, "tokens exchange information", ha="center", va="top", fontsize=9.5,
            color=C_GREY)
    ax.text(9.5, y - 0.15, "each token on its own", ha="center", va="top", fontsize=9.5,
            color=C_GREY)
    ax.text(3.8, y + h + 0.55, "residual: input added back", ha="center", fontsize=9, color=C_GREY)
    return save(fig, path)


def fig_dit(path, shapes, n_params, depth, dim):
    """The whole transformer, left to right, every shape from a real forward pass."""
    fig, ax = canvas(13, 3.1, (-0.2, 15.2), (-0.75, 3.0))
    n_tok, d = shapes["block1"]
    g = shapes["patchify"][1]
    y, h = 1.2, 1.1
    stages = [
        (0.0, 1.6, f"input $x_t$\n1 x 32 x 32", "#9aa3ad"),
        (2.1, 2.1, f"cut into patches\nLinear 4 -> {d}\n{n_tok} x {d}", C_COOL),
        (4.7, 1.6, f"+ position\ntable\n{n_tok} x {d}", C_COOL),
        (6.8, 2.6, f"{depth} x transformer block\n{shape_str(shapes['block_last'])}\n(size never changes)",
         "#3d4a5c"),
        (9.9, 2.0, f"norm, Linear\n{d} -> 4\n{shape_str(shapes['head'])}", C_GREY),
        (12.4, 2.6, f"put the patches back\noutput $\\hat\\epsilon$\n{shape_str(shapes['output'])}", C_WARM),
    ]
    for x, w, label, col in stages:
        box(ax, x, y, w, h, label, col, fs=10)
    for (x, w, *_), (nx, *_) in zip(stages, stages[1:]):
        arrow(ax, (x + w, y + h / 2), (nx, y + h / 2))
    ax.text(3.15, y + h + 0.2, f"{g} x {g} = {n_tok} patches", ha="center", fontsize=9, color=C_GREY)

    box(ax, 6.8, -0.65, 5.1, 0.5, f"e: {dim} numbers made from t and y (as in Part 1)",
        "#e8d6c8", fs=10, tc="#1c222b", alpha=1.0)
    for x in (8.1, 10.9):
        arrow(ax, (x, -0.15), (x, y), color=C_WARM, lw=1.2, ls=":")
    ax.text(12.1, -0.4, "as scale, shift, gate: see \"how t and y get in\"", fontsize=9.5, color=C_WARM, va="center")
    ax.set_title(f"week 5 transformer: {n_params:,} parameters. No pooling, no skips.",
                 fontsize=12, color="#1c222b")
    return save(fig, path)


def fig_loss_compare(path, log1, log2):
    """Both parts' training loss: per epoch (left) and per minute of training (right)."""
    import numpy as np
    fig, axes = plt.subplots(1, 2, figsize=(11, 3.7))
    fig.patch.set_facecolor(PAPER_HEX)
    for log, name, col, ls in ((log1, "U-Net (Part 1)", C_COOL, "--"),
                               (log2, "transformer (Part 2)", C_WARM, "-")):
        minutes = np.cumsum(log["seconds"]) / 60
        axes[0].plot(log["epoch"], log["loss"], marker="o", ms=4, color=col, ls=ls, lw=2, label=name)
        axes[1].plot(minutes, log["loss"], marker="o", ms=4, color=col, ls=ls, lw=2, label=name)
    style(axes[0], "epoch", "training loss (MSE)")
    style(axes[1], "minutes of training on this Mac", "training loss (MSE)")
    axes[0].set_xticks([e for e in log1["epoch"] if e == 1 or e % 5 == 0])
    axes[0].legend(frameon=False)
    return save(fig, path)


# ------------------------------------------------------------------ extra slide types

def slept_note(log):
    """The epochs whose time was not measured (the Mac slept), if any."""
    est = (log or {}).get("estimated_seconds")
    if not est:
        return ""
    return (f" Transformer epoch {', '.join(map(str, est))}: the Mac went to sleep mid-epoch, so its time is the "
            "average of the other epochs.")


def not_run_slide(prs, kicker, title, commands, what):
    """The honest placeholder: this slide's evidence does not exist yet."""
    s = blank(prs)
    header(s, kicker, title)
    bullets(s, [(f"Not run yet. {what} comes from:", "lead")]
            + [(c, "mono") for c in commands]
            + [("Then rebuild the deck:  python slides/build_slides.py", "note")])
    return s


def two_picture_slide(prs, kicker, title, top_img, top_h, bottom_img, bottom_h, cap=None):
    s = blank(prs)
    header(s, kicker, title)
    pic = picture(s, top_img, top=Inches(1.85), height=top_h)
    picture(s, bottom_img, top=Emu(int(pic.top) + int(pic.height) + Inches(0.1)), height=bottom_h)
    if cap:
        caption(s, cap, y=Inches(6.7))
    return s


def text_and_picture_slide(prs, kicker, title, items, img, img_top, img_h, cap=None, size=16):
    s = blank(prs)
    header(s, kicker, title)
    bullets(s, items, top=Inches(1.85), size=size)
    picture(s, img, top=img_top, height=img_h)
    if cap:
        caption(s, cap, y=Inches(6.75))
    return s


# ------------------------------------------------------------------ deck

def build():
    tmp = tempfile.mkdtemp()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    base = trained_base()
    model = UNet(base=base).eval()
    n_params = count_params(model)
    shapes = real_shapes(model)
    sched = NoiseSchedule()
    images, digits = load_mnist("cpu")
    x0 = images[(digits == 7).nonzero()[0]]                  # the first 7 in MNIST
    train_log = load_json("train_log.json")
    sample_log = load_json("sample_log.json") or {}
    emb = base * 4

    # ---------------------------------------------------------------- opening
    flow_slide(prs, "THE PROBLEM", "Generating has no single right answer", [
        ("lead", "What we want: new images of a digit we choose -- not copies of the training set."),
        ("text", "Weeks 3 and 4 had one correct answer per input (a class, a mask). "
                 "We compared the output to it: that was the loss."),
        ("text", "\"Draw a 7\" has millions of correct answers. There is nothing to compare the output to."),
        ("gap", 0.25),
        ("lead", "The idea: destroying an image is easy. Learn to undo it."),
        ("text", "Add noise, a little at a time, until only noise is left. We know exactly which "
                 "noise we added, so \"find the noise\" DOES have one right answer: an ordinary "
                 "supervised problem."),
        ("text", "To generate: start from pure noise and remove noise step by step until a digit is left."),
    ], tmp)

    # ---------------------------------------------------------------- forward process
    flow_slide(prs, "ADDING NOISE (THE FORWARD PROCESS)", "Any noise level in one formula", [
        ("eq", EQ_SCHEDULE, 24),
        ("note", "beta_t (beta): how much noise step t adds -- a fixed list of 1000 numbers, "
                 "rising from 0.0001 to 0.02.   alpha_t = 1 - beta_t: how much of the image that step keeps.   "
                 "alpha-bar_t: the product of all alphas up to t = how much of the original digit is left."),
        ("gap", 0.15),
        ("eq", EQ_FORWARD, 26),
        ("note", "x_0: the clean digit, pixels in [-1, 1].   epsilon: random noise, one number per pixel "
                 "drawn from a bell curve with mean 0 and spread 1.   x_t: the digit at noise level t."),
        ("gap", 0.15),
        ("text", "No loop: any level t is reached in one line. So training can pick a random t for "
                 "every image in the batch."),
    ], tmp)

    two_picture_slide(prs, "ADDING NOISE", "A real training digit at eight noise levels",
                      fig_forward(os.path.join(tmp, "fwd.png"), x0, sched), Inches(2.2),
                      fig_schedule(os.path.join(tmp, "sched.png"), sched), Inches(2.45),
                      cap="Top: x_t computed with the formula above.  Bottom: the two weights "
                          "in the formula. By t = 999 the digit is gone.")

    # ---------------------------------------------------------------- the model's job
    figure_slide(prs, "TRAINING", "One training step: where each input comes from, and what learns",
                 fig_training_flow(os.path.join(tmp, "train_flow.png"), x0, sched, emb),
                 height=Inches(4.75), top=Inches(1.8),
                 cap="The label table is a weight like any other: backprop changes row 7 whenever "
                     "the training digit is a 7.")

    two_picture_slide(prs, "TRAINING", "How t and y become e, and where e goes",
                      fig_embedding(os.path.join(tmp, "embed.png"), emb), Inches(2.55),
                      fig_block(os.path.join(tmp, "block.png"), shapes, base), Inches(1.95),
                      cap="Top: t -> sin/cos (fixed) -> MLP (learned); y -> its row of the table "
                          "(learned). Their sum is e.  Bottom: inside every block, a learned Linear "
                          "turns e into one number per channel, added at every pixel.")

    flow_slide(prs, "THE LOSS", "Predict the noise; the error is plain MSE", [
        ("eq", EQ_LOSS, 28),
        ("note", "epsilon_theta: the U-Net (theta = its weights), taking the noisy digit, the level "
                 "and the label.   E[...]: the average over digits x_0, labels y, random levels t and "
                 "random noise epsilon.   ||...||^2: squared error over all 32 x 32 pixels -- MSE."),
        ("gap", 0.2),
        ("text", "Why the noise and not the clean digit? The next two slides."),
    ], tmp)

    flow_slide(prs, "WHY PREDICT THE NOISE (1/2)", "Knowing the noise = knowing the digit", [
        ("eq", EQ_FORWARD, 24),
        ("text", "The model already has x_t (its input) and t. That is one equation with two "
                 "unknowns, x_0 and epsilon. Know one, and the other follows:"),
        ("eq", EQ_X0_FROM_EPS, 24),
        ("eq", EQ_EPS_FROM_X0, 24),
        ("text", "So a model that outputs the noise has also told us its guess for the digit. The two "
                 "choices do not differ in what the model knows. They differ in what the loss punishes."),
    ], tmp)

    s = blank(prs)
    header(s, "WHY PREDICT THE NOISE (2/2)", "What the loss cares about")
    math_block(s, EQ_WEIGHT, M, Inches(1.8), tmp, fontsize=22, center=True)
    picture(s, fig_loss_weight(os.path.join(tmp, "weight.png"), x0, sched),
            top=Inches(2.55), height=Inches(2.75))
    bullets(s, [
        "Predict the noise: a small mistake on a NEARLY CLEAN image costs a lot (left of the curve). "
        "The network spends its effort on fine details -- stroke edges, thickness, clean background.",
        "Predict the digit: the same small mistake costs almost nothing. The loss is dominated by very "
        "noisy images, where the digit must be guessed from almost pure noise -- the best possible "
        "answer is a blurry average, and the loss keeps pushing on a task that cannot be solved.",
        ("Predicting the noise makes the loss care most about the details that decide how the final "
         "digit looks.", "lead"),
    ], top=Inches(5.35), size=13.5)

    code_slide(prs, "TRAINING", "One training step: the whole algorithm", "part1_unet_diffusion.py", [
        "t = torch.randint(0, sched.T, (len(x0),), device=device)   # a random level per image",
        "noise = torch.randn_like(x0)                                # random noise",
        "x_t = sched.add_noise(x0, t, noise)                         # the forward formula",
        ("loss = F.mse_loss(model(x_t, t, y), noise)                 # find the noise",),
        "",
        "opt.zero_grad(set_to_none=True)",
        "loss.backward()",
        "opt.step()",
    ], note="Two extras, a few lines each: label dropout (10% of labels replaced by \"no label\", "
            "used later for guidance) and EMA, a second copy of the weights that is a slowly moving "
            "average of the trained ones. Its samples are cleaner; the checkpoint stores it.")

    if train_log:
        last = train_log["epoch"][-1]
        grids = [(p, f"after epoch {e}") for e in sorted({1, (last + 1) // 2, last})
                 if (p := result(f"epoch_{e:02d}.png"))]
        s = blank(prs)
        header(s, "RESULTS: TRAINING",
               "The loss falls, and digits appear" if grids else "The loss falls, then flattens")
        summary = (f"{last} epochs, {sum(train_log['seconds']) / 60:.0f} min on this Mac, "
                   f"{train_log['params']:,} parameters.")
        live = " Live: trackio show --project week5-diffusion"
        if grids:
            pic = s.shapes.add_picture(fig_loss(os.path.join(tmp, "loss.png"), train_log),
                                       M, Inches(1.95), height=Inches(3.9))
            g = s.shapes.add_picture(fig_images(os.path.join(tmp, "grids.png"), grids),
                                     Emu(int(pic.left) + int(pic.width) + Inches(0.2)),
                                     Inches(1.95), height=Inches(3.9))
            room = int(W - Inches(0.3)) - int(g.left)
            if int(g.width) > room:                  # too wide: shrink, keep proportions
                g.height, g.width = int(g.height * room / int(g.width)), room
            caption(s, summary + " Right: rows are digits 0-9, same starting noise after every "
                    "epoch (made with a faster 50-step sampler with guidance, not covered here)."
                    + live, y=Inches(6.2))
        else:                                        # no per-epoch grids: the loss curve alone
            pic = s.shapes.add_picture(fig_loss(os.path.join(tmp, "loss.png"), train_log),
                                       M, Inches(1.95), height=Inches(4.1))
            pic.left = int((W - pic.width) / 2)
            caption(s, summary + live, y=Inches(6.2))
    else:
        not_run_slide(prs, "RESULTS: TRAINING", "The loss falls, and digits appear",
                      ["python part1_unet_diffusion.py"], "The loss curve and the per-epoch sample grids")

    figure_slide(prs, "INFERENCE", "Generating a 7: same network, run in a loop, nothing learns",
                 fig_inference_flow(os.path.join(tmp, "infer_flow.png"), emb),
                 height=Inches(4.75), top=Inches(1.8),
                 cap="The seed decides WHICH 7 (slant, thickness). The label decides THAT it is a 7. "
                     "The model always runs forward; only t counts down.")

    if (p := result("trajectory_ddpm.png")):
        figure_slide(prs, "RESULTS: INFERENCE", "Noise on the left, digit on the right",
                     p, height=Inches(4.4),
                     cap="One row per requested digit, 0-9. Columns: 10 evenly spaced moments of the "
                         "1000-step loop from the previous slide.")
    else:
        not_run_slide(prs, "RESULTS: INFERENCE", "Noise on the left, digit on the right",
                      ["python sample.py --sampler ddpm --guidance 1 --trajectory"],
                      "The step-by-step picture")

    if (p := result("samples_ddpm_w1.png")):
        timing = sample_log.get("ddpm_1000steps_w1", {})
        extra = (f" {timing['digits']} digits x 1000 steps took {timing['seconds']:.0f} s."
                 if timing else "")
        figure_slide(prs, "RESULTS: INFERENCE", "Ask for each digit, get eight new ones",
                     p, height=Inches(4.4),
                     cap="Row k: y = k, eight different seeds. Same trained model, 1000-step loop, "
                         "nothing else." + extra)

    part2(prs, tmp, x0, model, train_log)

    prs.save(OUT)
    print(f"wrote {OUT}  ({len(prs.slides)} slides)")


def part2(prs, tmp, x0, unet, train_log):
    """Part 2: the same diffusion model with a transformer as the network."""
    dit = load_dit()
    d_params, u_params = count_params(dit), count_params(unet)
    shapes = dit_shapes(dit)
    n_tok, dim = shapes["block1"]
    log2 = load_json("train_log.json", RESULTS2)
    planned = 20

    section_slide(prs, "PART 2", "Swap the U-Net for a transformer",
                  "Same noise, same loss, same training loop, same samplers, same e. Only the "
                  f"network that predicts the noise changes -- sized to the same parameter count: "
                  f"{d_params:,} vs the U-Net's {u_params:,}.")

    # ------------------------------------------------ from an image to tokens
    figure_slide(prs, "PART 2: THE MODEL", "From an image to tokens",
                 fig_patches(os.path.join(tmp, "patches.png"), dit, x0),
                 height=Inches(4.1), top=Inches(1.95),
                 cap="A token is one patch, described by numbers the network works with. The same Linear "
                     "is used for every patch. The position table gives each patch its own learned row, "
                     "so the network knows where the patch sits.")

    # ------------------------------------------------ attention, on its own
    toy_h, *_, toy_a, _ = toy_attention()
    toy_terms = " + ".join(f"{w:.2f} v{j + 1}" for j, w in enumerate(toy_a[2].tolist()))
    flow_slide(prs, "PART 2: ATTENTION (1/2)", "Every token takes a weighted mix of all tokens", [
        ("lead", "A transformer works on tokens: rows of d numbers, each describing one piece of the "
                 f"input (ours: one 2 x 2 patch of the image, so N = {n_tok} and d = {dim})."),
        ("text", "In attention, each token asks a question, every token offers an answer, and how well "
                 "they match decides how much of each answer to take."),
        ("eq", EQ_QKV, 24),
        ("note", "h: the N tokens, an N x d table.   W_Q, W_K, W_V: learned d x d matrices, the same for "
                 "every token.   Q (query): what each token is looking for.   K (key): what each token "
                 "contains.   V (value): what each token passes on when it is chosen."),
        ("eq", EQ_ATTN, 26),
        ("note", "Q K^T: an N x N table of match scores (dot products), one per pair of tokens.   "
                 "/ sqrt(d): stops the scores growing with d, so softmax does not put all the weight on "
                 "one token.   softmax: each row becomes positive weights that sum to 1.   x V: each "
                 "token's output is a weighted average of all the tokens' values."),
        ("text", "Unlike a convolution's fixed 3 x 3 window, attention decides from the input which "
                 "tokens to mix."),
    ], tmp)

    figure_slide(prs, "PART 2: ATTENTION (2/2)",
                 f"The same formula, on {len(toy_h)} tokens of {toy_h.shape[1]} numbers",
                 fig_attention_toy(os.path.join(tmp, "attn_toy.png")),
                 height=Inches(4.55), top=Inches(1.8),
                 cap=f"Toy numbers: random tokens and random W_Q, W_K, W_V (seed 0). Token 3 (orange): "
                     f"output = {toy_terms}  (v_j = row j of V). Each row of weights is different: "
                     "every token decides for itself whom to listen to.")

    flow_slide(prs, "PART 2: WHY ATTENTION", "What attention offers that a convolution does not", [
        ("text", "1.  The whole image in one step. A 3 x 3 convolution sees only its neighbours; to see "
                 "far, a CNN needs many layers or must shrink the image (the U-Net goes down to 4 x 4). "
                 "Attention lets every patch use every other patch, in every block."),
        ("text", "2.  The mixing is chosen by the input. A convolution applies the same weights to every "
                 "image. Attention computes its weights from the tokens themselves (previous slide)."),
        ("text", "3.  Few built-in assumptions, so it keeps improving with more data. A convolution "
                 "assumes nearby pixels matter most, which helps when data is scarce. Attention must "
                 "learn even which patches are neighbours (the position table), so it needs more data, "
                 "but given a lot of data and a bigger model it keeps getting better."),
        ("text", "4.  One design for every kind of input. Words, image patches and sound all become "
                 "tokens, and tokens of different kinds can attend to each other: image patches can "
                 "look at the words of a prompt."),
        ("text", "5.  Mostly large matrix multiplications, which GPUs are built for."),
        ("note", f"The price: every token is scored against every token. Our {n_tok} tokens give a "
                 f"{n_tok} x {n_tok} table of {n_tok * n_tok:,} scores; one token per pixel would give "
                 f"1,024 x 1,024 = {1024 * 1024:,}. That is why the image is cut into 2 x 2 patches. "
                 "On a dataset as small as ours, point 3 works against it: see the results."),
    ], tmp)

    # ------------------------------------------------ the transformer that uses it
    s = blank(prs)
    header(s, "PART 2: THE MODEL", "One transformer block: attention, then an MLP")
    picture(s, fig_plain_block(os.path.join(tmp, "plain_block.png"), n_tok, dim, len(dit.blocks)),
            top=Inches(1.95), height=Inches(2.3))
    bullets(s, [
        f"LN (LayerNorm): rescales each token's {dim} numbers to average 0, spread 1, so every "
        "block gets inputs in the same range.",
        "Attention: the tokens exchange information (attention slides).   MLP: two Linear layers with "
        "a GELU between them, applied to each token on its own, the same weights for all tokens.",
        "Residual: each half's output is added onto its input, so a block only has to learn a change, "
        "and the gradient reaches early blocks directly.",
        "Shape in = shape out, so blocks stack: the whole network is this block, repeated.",
    ], top=Inches(4.45), size=14.5)

    g = IMAGE_SIZE // dit.patch
    ex = 74
    figure_slide(prs, "PART 2: THE MODEL", "Image in, image out: tokens never change order",
                 fig_dit(os.path.join(tmp, "dit.png"), shapes, d_params, len(dit.blocks), dim),
                 height=Inches(3.3), top=Inches(2.1),
                 cap=f"Row i is always patch i: no block reorders the tokens. So at the end, Linear "
                     f"{dim} -> 4 turns each token into the predicted noise for its own 2 x 2 pixels, and "
                     f"each patch goes back to its place (token {ex}: grid row {ex // g}, column {ex % g}, "
                     f"so pixel rows {2 * (ex // g)}-{2 * (ex // g) + 1}, columns "
                     f"{2 * (ex % g)}-{2 * (ex % g) + 1}). Unlike the U-Net, nothing shrinks: all "
                     f"{n_tok} patches stay at full detail in every block.")

    pos_png, pa, pb = fig_position(os.path.join(tmp, "position.png"), dit, x0)
    _, near, far = position_stats(dit)
    figure_slide(prs, "PART 2: THE MODEL", "The position table: a learned address for each patch",
                 pos_png, height=Inches(4.0), top=Inches(1.85),
                 cap=f"Attention compares only what tokens contain, never where they sit. A and B contain "
                     f"the same 4 numbers, so without the table they would be the same token. Adding row "
                     f"{pa} vs row {pb} makes them differ. Nobody told the table which patches are "
                     f"neighbours: after training, rows of neighbouring patches have similarity "
                     f"{near:.2f} on average, rows 6 or more grid steps apart {far:.2f}.")

    u_gf, d_gf = gflops(unet), gflops(dit)
    rows = [["parameters", f"{u_params:,}", f"{d_params:,}"],
            ["compute per image, forward (GFLOPs)", f"{u_gf:.2f}", (f"{d_gf:.2f}",)]]
    u_sec = sum(train_log["seconds"]) / len(train_log["seconds"]) if train_log else None
    d_sec = sum(log2["seconds"]) / len(log2["seconds"]) if log2 else None
    if u_sec and d_sec:
        rows.append(["seconds per training epoch (average)", f"{u_sec:.0f}", (f"{d_sec:.0f}",)])
    if train_log and log2:
        rows.append([f"minutes for {planned} epochs",
                     f"{sum(train_log['seconds']) / 60:.0f}",
                     (f"{d_sec * planned / 60:.0f}" + ("" if len(log2["epoch"]) >= planned
                                                         else " (projected)"),)])
    table_slide(prs, "PART 2: THE COST", "Same parameters, not the same compute",
                ["", "U-Net (Part 1)", "transformer (Part 2)"], rows, widths=[2.2, 1.2, 1.4],
                note="GFLOPs: billions of multiplications and additions to process one image once, "
                     "counted by PyTorch. Every transformer weight is used on all "
                     f"{n_tok} tokens; the U-Net's widest layers only run on 4 x 4 = 16 positions. "
                     "Attention is also slower per GFLOP than convolution on this Mac's GPU."
                     + slept_note(log2))

    if log2:
        done = len(log2["epoch"])
        status = (f"{done} epochs" if done >= planned
                  else f"{done} of {planned} epochs so far (training still running)")
        if log2.get("resumed_after"):
            r = log2["resumed_after"]
            status += (f" (stopped after epoch {r}, then continued from the saved weights: the small "
                       f"bump at epoch {r + 1})")
        figure_slide(prs, "PART 2: RESULTS", "Training loss: per epoch, and per minute",
                     fig_loss_compare(os.path.join(tmp, "loss_compare.png"), train_log, log2),
                     height=Inches(3.9), top=Inches(1.95),
                     cap=f"Transformer: {status}, final loss {log2['loss'][-1]:.4f}. U-Net: "
                         f"{len(train_log['epoch'])} epochs, final loss {train_log['loss'][-1]:.4f}. "
                         "Same data, same loss, same recipe, so the numbers compare directly."
                         + slept_note(log2))
        last = log2["epoch"][-1]
        grids = [(p, f"after epoch {e}") for e in sorted({1, (last + 1) // 2, last})
                 if (p := result(f"epoch_{e:02d}.png", RESULTS2))]
        if grids:
            figure_slide(prs, "PART 2: RESULTS", "The transformer's digits, epoch by epoch",
                         fig_images(os.path.join(tmp, "grids2.png"), grids),
                         height=Inches(4.1), top=Inches(1.95),
                         cap="Rows are digits 0-9, same starting noise after every epoch "
                             "(50-step sampler with guidance w = 3).")
    else:
        not_run_slide(prs, "PART 2: RESULTS", "Training loss: per epoch, and per minute",
                      ["python part2_transformer_diffusion.py"], "The loss curve")


if __name__ == "__main__":
    build()
